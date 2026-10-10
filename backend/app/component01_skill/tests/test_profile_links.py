import io
from unittest.mock import patch
import httpx
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import pymupdf
from app.component01_skill.services.cv_parser import parse_cv, parse_cv_document
from app.component01_skill.services.profile_identifier import detect_identifiers, normalize_identifier
from app.component01_skill.services.github_analyzer import analyze_github, _public_metadata
from app.component01_skill.services.linkedin_evidence import analyze_linkedin, register_authorized_provider
from app.component01_skill.services.profile_service import process_cv


def linked_pdf():
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((40, 40), 'SKILLS\nPython')
        # Rectangle/icon region has no visible URL or text.
        page.draw_rect(pymupdf.Rect(40, 80, 60, 100))
        for rect, url in [(pymupdf.Rect(40, 80, 60, 100), 'https://github.com/Example'),
                          (pymupdf.Rect(80, 80, 100, 100), 'https://www.linkedin.com/in/example-person')]:
            page.insert_link({'kind': pymupdf.LINK_URI, 'from': rect, 'uri': url})
        return doc.tobytes()


def test_pdf_hidden_links_are_not_skill_text():
    text, links = parse_cv_document('cv.pdf', linked_pdf())
    assert text == parse_cv('cv.pdf', linked_pdf())
    assert 'github' not in text.lower()
    assert len(detect_identifiers(text, links=links)[0]) == 2


def test_docx_text_icon_and_header_relationships():
    doc = Document()
    doc.add_paragraph('SKILLS')
    doc.add_paragraph('Python')
    paragraph = doc.add_paragraph()
    rid = doc.part.relate_to('https://github.com/Example', RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = 'My code'
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    # Drawing link relationships use the same hyperlink relationship type.
    icon_rid = doc.part.relate_to('https://www.linkedin.com/in/example-person', RT.HYPERLINK, is_external=True)
    drawing = OxmlElement('w:drawing')
    click = OxmlElement('a:hlinkClick')
    click.set(qn('r:id'), icon_rid)
    drawing.append(click)
    paragraph._p.append(drawing)
    doc.sections[0].header.part.relate_to('https://github.com/HeaderUser', RT.HYPERLINK, is_external=True)
    stream = io.BytesIO()
    doc.save(stream)
    visible, links = parse_cv_document('cv.docx', stream.getvalue())
    assert 'My code' in visible
    assert len(links) == 3
    assert len(detect_identifiers(visible, links=links)[0]) == 3


@pytest.mark.parametrize('url', [
    'https://github.com.evil.test/user', 'https://github.com@evil.test/user',
    'https://evil.test/github.com/user', 'file://github.com/user',
    'https://github.com:443/user', 'https://github.com/user/repo',
    'https://github.com/user%2Fother', 'https://github.com/user\\evil',
    'https://github.com/login', 'https://github.com/user--name',
])
def test_invalid_urls_rejected_before_network(url):
    assert normalize_identifier('github', url) is None
    with patch('httpx.Client', side_effect=AssertionError('No request allowed')):
        assert analyze_github({'url': url})[1] == 'invalid'


def test_multiple_profiles_dedup_normalization_and_plain_handles():
    profiles, warnings = detect_identifiers(
        'GitHub username: @Example\nLinked In: example-person\nhttps://github.com/EXAMPLE/?tab=repositories\n'
        'github.com/Second\nhttps://github.com.evil.test/trap',
        links=['https://github.com/example', 'http://www.linkedin.com/in/example-person/?tracking=x'])
    assert [(p['platform'], p['handle']) for p in profiles] == [('github', 'Example'), ('github', 'Second'), ('linkedin', 'example-person')]
    assert profiles[2]['url'] == 'https://www.linkedin.com/in/example-person'
    assert warnings


def test_github_validation_repository_filtering_and_no_redirects():
    _public_metadata.cache_clear()
    requests = []
    def handler(request):
        requests.append(request)
        assert request.url.host == 'api.github.com'
        if request.url.path.endswith('/repos'):
            return httpx.Response(200, json=[
                {'name': 'project', 'language': 'Python', 'owner': {'login': 'Example'}},
                {'name': 'fork', 'language': 'Python', 'fork': True, 'owner': {'login': 'Example'}},
                {'name': 'foreign', 'language': 'Python', 'owner': {'login': 'other'}}])
        return httpx.Response(200, json={'login': 'Example', 'type': 'User'})
    real_client = httpx.Client
    with patch('app.component01_skill.services.github_analyzer.collect_usage', return_value=([], [])), patch('httpx.Client', side_effect=lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw)):
        records, status, warnings = analyze_github({'url': 'https://github.com/Example'})
    assert status == 'completed' and len(records) == 1 and len(requests) == 2
    assert records[0]['skill'] == 'Python'
    assert warnings
    _public_metadata.cache_clear()


@pytest.mark.parametrize('status', [301, 403, 404, 429, 500])
def test_github_inaccessible_profiles_are_optional(status):
    _public_metadata.cache_clear()
    real_client = httpx.Client
    with patch('httpx.Client', side_effect=lambda **kw: real_client(transport=httpx.MockTransport(lambda req: httpx.Response(status)), **kw)):
        assert analyze_github({'url': 'https://github.com/Example'})[1] == 'unavailable'


def test_linkedin_authorized_adapter_and_failure():
    profile = {'url': 'https://linkedin.com/in/example-person'}
    assert analyze_linkedin('candidate', profile)[1] == 'not_performed'
    calls = []
    try:
        register_authorized_provider(lambda cid, url: calls.append((cid, url)) or [])
        assert analyze_linkedin('candidate', profile)[1] == 'completed'
        assert calls == [('candidate', 'https://www.linkedin.com/in/example-person')]
        register_authorized_provider(lambda *args: (_ for _ in ()).throw(PermissionError()))
        assert analyze_linkedin('candidate', profile)[1] == 'unavailable'
    finally:
        register_authorized_provider(None)


def test_upload_automatic_analysis_dedup_and_unchanged_cv_levels(monkeypatch):
    monkeypatch.setenv('SKILL_ENABLE_SBERT', 'false')
    record = {'skill': 'Python', 'text': 'Built Python production systems.', 'source': 'github', 'url': 'https://github.com/Example/project'}
    with patch('app.component01_skill.services.profile_service.repository.save_assessment', side_effect=lambda cid, content, details: details), \
         patch('app.component01_skill.services.profile_service.analyze_github_profile', return_value=([record, record], 'completed', [])) as github, \
         patch('app.component01_skill.services.profile_service.analyze_linkedin', return_value=([dict(record, source='linkedin')], 'completed', [])):
        baseline = process_cv('candidate', 'cv.pdf', linked_pdf())
        github.assert_called_once()
        enriched = process_cv('candidate', 'cv.pdf', linked_pdf(), analyze_github=True)
        assert github.call_count == 2
    assert set(enriched) == set(baseline)
    assert enriched['skills'][0]['level'] == baseline['skills'][0]['level'] == 2
    assert len(enriched['skills'][0]['evidence']) == 2
    assert enriched['external_profiles'][0]['ownership_verified'] is False


def test_same_project_across_cv_and_profiles_is_not_counted_twice():
    from app.component01_skill.services.profile_evidence import combine_evidence
    skills = [{'name': 'Python', 'level': 3, 'evidence': [
        {'section': 'projects', 'text': 'Built Python tools for Weather Dashboard.'}]}]
    records = [{'skill': 'Python', 'source': 'github', 'project_name': 'weather-dashboard',
                'text': 'Repository uses Python.', 'url': 'https://github.com/Example/weather-dashboard'},
               {'skill': 'React', 'source': 'github', 'text': 'React repository.'}]
    result = combine_evidence(skills, records)
    assert result == skills and len(result[0]['evidence']) == 1
    assert result[0]['level'] == 3


def test_github_timeout_and_invalid_response_do_not_block():
    for response in ('timeout', 'malformed'):
        _public_metadata.cache_clear()
        def handler(request):
            if response == 'timeout':
                raise httpx.ReadTimeout('timeout', request=request)
            return httpx.Response(200, content=b'not json')
        real_client = httpx.Client
        with patch('httpx.Client', side_effect=lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw)):
            assert analyze_github({'url': 'https://github.com/Example'})[1] == 'unavailable'
