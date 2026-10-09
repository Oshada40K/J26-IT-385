"""Regression cases for arbitrary CVs, including the supplied fictional sample."""
import io
import json
import os
import zipfile
from pathlib import Path
from unittest.mock import patch
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from app.component01_skill.router import router
from app.component01_skill.services.cv_parser import CVError, parse_cv
from app.component01_skill.services.profile_identifier import detect_identifiers
from app.component01_skill.services.skill_extractor import extract_skills
from app.component01_skill.services.skill_assessor import assess_skills

SAMPLE = Path(__file__).parents[2] / 'app/component01_skill/Tharindi_Jayawickrama_Sample_CV_Usernames_Only.pdf'


def pdf(text):
    import pymupdf
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((40, 40), text)
        return doc.tobytes()


def assess(text):
    with patch.dict(os.environ, {'SKILL_ENABLE_SBERT': 'false'}):
        return {item['name']: item for item in assess_skills(extract_skills(text)[0])}


def test_aliases_boundaries_and_unmentioned_skills():
    result = assess('TECHNICAL SKILLS\nReact, ReactJS, React.js, JavaScript, Python\n')
    assert set(result) == {'React', 'JavaScript', 'Python'}
    assert result['Python']['level'] == 2
    assert result['Python']['confidence'] == 'low'
    assert result['React']['status'] == 'provisional'


def test_negative_and_speculative_contexts():
    result = assess('SUMMARY\nNo Python experience. I want to learn React.\nI react to challenges and go to conferences.\nPROJECTS\nBuilt a JavaScript dashboard.\n')
    assert set(result) == {'JavaScript'}
    assert result['JavaScript']['level'] == 3


def test_projects_link_technology_stack_to_actions():
    result = assess('PROJECTS\n01 Web dashboard\nReact | PostgreSQL\nBuilt and tested a dashboard.\n02 Test suite\nPython | Pytest\nDesigned automated browser tests.\n')
    assert result['React']['level'] == 3
    assert result['Python']['level'] == 3
    spaced = assess('PROJECTS\n\n01 Dashboard\n\nReact | PostgreSQL\n\nBuilt a dashboard.\n\n02 QA Suite\n\nPython | Selenium\n\nDesigned automated tests.\n')
    assert spaced['Python']['level'] == 3
    assert 'Python' not in spaced['React']['evidence'][0]['text']
    assert 'Python' not in result['React']['evidence'][0]['text']


def test_deduplicated_evidence_and_bounded_levels():
    result = assess('EXPERIENCE\nBuilt React dashboards.\nBuilt React dashboards.\n')
    assert len(result['React']['evidence']) == 1
    assert all(1 <= item['level'] <= 5 for item in result.values())
    assert result['React']['level'] != 5
    assert assess('EDUCATION\nIntroduction to Python programming.\n')['Python']['level'] == 1
    assert assess('SUMMARY\nFriendly team player with good communication.\n') == {}


def test_identifiers_and_missing_profiles():
    profiles, warnings = detect_identifiers('GitHub: TharindiJay2002\nLinkedIn: tharindi-jayawickrama-36314a366')
    assert len(profiles) == 2 and not warnings
    assert profiles[0]['url'] == 'https://github.com/TharindiJay2002'
    assert profiles[1]['ownership_verified'] is False
    profiles, warnings = detect_identifiers('', 'https://evil.test/name')
    assert not profiles and warnings
    assert detect_identifiers('')[0] == []


def test_sample_pdf_is_real_input_not_fixture_scores():
    text = parse_cv(SAMPLE.name, SAMPLE.read_bytes())
    result = assess(text)
    assert {'React', 'Python', 'Selenium', 'PostgreSQL', 'MongoDB', 'Postman'} <= set(result)
    assert result['React']['level'] == 3
    assert result['Python']['level'] == 3
    assert result['Selenium']['level'] == 3
    assert result['Node.js']['level'] == 3
    assert result['TypeScript']['level'] == 2
    assert result['React']['evidence']
    expected = json.loads((Path(__file__).parent / 'fixtures/sample_expected.json').read_text())['technical_skills']
    assert [{'name': item['name'], 'level': item['level']} for item in result.values()] == expected
    assert len(detect_identifiers(text)[0]) == 2


def test_docx_table_and_paragraph_order():
    from docx import Document
    doc = Document()
    doc.add_paragraph('TECHNICAL SKILLS')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Python'
    table.cell(0, 1).text = 'SQL'
    stream = io.BytesIO()
    doc.save(stream)
    assert set(assess(parse_cv('resume.DOCX', stream.getvalue()))) == {'Python', 'SQL'}


def test_bad_empty_scanned_encrypted_and_oversized_inputs():
    with pytest.raises(CVError): parse_cv('cv.pdf', b'invalid')
    with pytest.raises(CVError): parse_cv('cv.exe', b'anything')
    with pytest.raises(CVError): parse_cv('cv.docx', b'PK corrupt')
    with pytest.raises(CVError): parse_cv('cv.pdf', b'')
    with pytest.raises(CVError): parse_cv('cv.pdf', b'x' * (10 * 1024 * 1024 + 1))
    import pymupdf
    with pymupdf.open() as doc:
        doc.new_page()
        blank = doc.tobytes()
        encrypted = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw='secret', owner_pw='owner')
    with pytest.raises(CVError, match='OCR'): parse_cv('scan.pdf', blank)
    with pytest.raises(CVError, match='Password'): parse_cv('encrypted.pdf', encrypted)
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', 'x' * (26 * 1024 * 1024))
    with pytest.raises(CVError, match='expanded'): parse_cv('bomb.docx', stream.getvalue())


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('SKILL_DATABASE_PATH', str(tmp_path / 'skills.sqlite3'))
    monkeypatch.setenv('SKILL_ENABLE_SBERT', 'false')
    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as client:
        yield client


def candidate(client):
    result = client.post('/api/component01/skills/session')
    assert result.status_code == 201
    session = result.json()
    return session['candidate_id'], {'Authorization': 'Bearer ' + session['access_token']}


def upload(client, candidate_id, headers, content, filename='cv.pdf', **extra):
    return client.post('/api/component01/skills/upload', data={'candidate_id': candidate_id, **extra}, files={'file': (filename, content)}, headers=headers)


def test_api_persistence_contract_history_and_authorization(client):
    cid, headers = candidate(client)
    path = '/api/component01/skills/' + cid
    assert client.get(path, headers=headers).status_code == 404
    assert client.get(path).status_code == 401
    first = upload(client, cid, headers, SAMPLE.read_bytes())
    assert first.status_code == 200, first.text
    assert set(first.json()) == {'candidate_id', 'assessment_id', 'status'}
    response = client.get(path, headers=headers).json()
    assert set(response) == {'schema_version', 'candidate_id', 'technical_skills'}
    assert response['schema_version'] == '1.0'
    assert all(set(skill) == {'name', 'level'} for skill in response['technical_skills'])
    other, other_headers = candidate(client)
    assert client.get(path, headers=other_headers).status_code == 403
    assert upload(client, cid, other_headers, SAMPLE.read_bytes()).status_code == 403
    assert upload(client, cid, headers, b'broken').status_code == 422
    assert client.get(path, headers=headers).json() == response
    with patch('app.component01_skill.router.process_cv', side_effect=AssertionError('GET must not infer')):
        assert client.get(path + '/details', headers=headers).status_code == 200
    second = upload(client, cid, headers, pdf('TECHNICAL SKILLS\nRust'))
    assert second.status_code == 200
    assert second.json()['assessment_id'] != first.json()['assessment_id']
    assert client.get(path, headers=headers).json()['technical_skills'] == [{'name': 'Rust', 'level': 2}]
    from app.component01_skill.repository import connection
    with connection() as db:
        assert db.execute('SELECT count(*) FROM skill_assessments WHERE candidate_id=?', (cid,)).fetchone()[0] == 2


def test_empty_skills_and_unavailable_enrichment(client):
    cid, headers = candidate(client)
    response = upload(client, cid, headers, pdf('SUMMARY\nTeam player.'), analyze_github='true', github_username='invalid/name')
    assert response.status_code == 200
    profile = client.get('/api/component01/skills/' + cid, headers=headers).json()
    assert profile['technical_skills'] == []
    details = client.get('/api/component01/skills/' + cid + '/details', headers=headers).json()
    assert details['warnings']


def test_sbert_cached_and_threshold_guarded(monkeypatch):
    import numpy as np
    from app.component01_skill.services import skill_extractor
    class FakeModel:
        def encode(self, sentences, **kwargs):
            return np.array([[1., 0.]] * len(sentences))
    monkeypatch.setenv('SKILL_ENABLE_SBERT', 'true')
    monkeypatch.setattr(skill_extractor, 'semantic_model', lambda: FakeModel())
    extracted, version, warnings = extract_skills('PROJECTS\nBuilt React dashboards.\n')
    assert set(extracted) == {'React'}
    assert extracted['React'][0]['semantic_supported'] is True
    assert assess_skills(extracted)[0]['level'] == 3
    monkeypatch.setattr(skill_extractor, 'semantic_model', lambda: (_ for _ in ()).throw(RuntimeError('unavailable')))
    extracted, version, warnings = extract_skills('SKILLS\nPython\n')
    assert extracted and any('unavailable' in value for value in warnings)


def test_short_aliases_are_not_nested_technologies():
    result = assess('SKILLS\nNode.js, Express.js, C++, C#\n')
    assert set(result) == {'Node.js', 'Express.js', 'C++', 'C#'}


def test_project_titles_without_numbers_and_pdf_spacing():
    result = assess('PROJECTS\nWeb Dashboard\n\nReact | PostgreSQL\n\nBuilt a dashboard.\n\nQA Automation Suite\n\nPython | Selenium\n\nDesigned automated tests.\n')
    assert result['Python']['level'] == 3
    assert result['React']['level'] == 3
    assert 'Python' not in result['React']['evidence'][0]['text']


def test_docx_api_and_size_limits(client):
    from docx import Document
    doc = Document()
    doc.add_paragraph('SKILLS')
    doc.add_paragraph('Ruby')
    stream = io.BytesIO()
    doc.save(stream)
    cid, headers = candidate(client)
    assert upload(client, cid, headers, stream.getvalue(), 'resume.docx').status_code == 200
    path = '/api/component01/skills/' + cid
    assert client.get(path, headers=headers).json()['technical_skills'] == [{'name': 'Ruby', 'level': 2}]
    assert upload(client, cid, headers, b'x' * (10 * 1024 * 1024 + 1)).status_code == 413
    assert client.get(path, headers=headers).json()['technical_skills'][0]['name'] == 'Ruby'
